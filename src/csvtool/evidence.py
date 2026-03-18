"""Word document evidence generator."""

import logging
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

from csvtool.models import ExecutionSummary, TestResult, StepResult

logger = logging.getLogger(__name__)


class EvidenceGenerator:
    """Generates Word document evidence reports."""

    def __init__(self, output_dir: Path):
        """Initialize generator with output directory."""
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def generate(self, summary: ExecutionSummary) -> Path:
        """Generate evidence document from execution summary."""
        doc = Document()

        # Set up styles
        self._setup_styles(doc)

        # Add header
        self._add_header(doc, summary)

        # Add summary table
        self._add_summary_table(doc, summary)

        # Add each test result
        for result in summary.test_results:
            self._add_test_result(doc, result)

        # Save document
        filename = self._generate_filename(summary)
        doc_path = self.output_dir / filename
        doc.save(str(doc_path))

        return doc_path

    def _setup_styles(self, doc: Document) -> None:
        """Set up document styles."""
        # Modify Normal style
        style = doc.styles["Normal"]
        style.font.name = "Arial"
        style.font.size = Pt(11)

    def _add_header(self, doc: Document, summary: ExecutionSummary) -> None:
        """Add document header."""
        # Title
        title = doc.add_heading("CSV TEST EVIDENCE REPORT", 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Metadata
        doc.add_paragraph()
        meta = doc.add_paragraph()
        meta.add_run(f"Application URL: ").bold = True
        meta.add_run(summary.url)

        meta = doc.add_paragraph()
        meta.add_run(f"Workbook: ").bold = True
        meta.add_run(summary.workbook_name)

        meta = doc.add_paragraph()
        meta.add_run(f"Execution Date: ").bold = True
        meta.add_run(summary.execution_date.strftime("%Y-%m-%d %H:%M:%S"))

        doc.add_paragraph()

    def _add_summary_table(self, doc: Document, summary: ExecutionSummary) -> None:
        """Add summary statistics table."""
        doc.add_heading("Execution Summary", level=1)

        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"

        # Header row
        header_cells = table.rows[0].cells
        header_cells[0].text = "Total Tests"
        header_cells[1].text = "Passed"
        header_cells[2].text = "Failed"
        header_cells[3].text = "Success Rate"

        # Make headers bold
        for cell in header_cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True

        # Data row
        row = table.add_row().cells
        row[0].text = str(summary.total_tests)
        row[1].text = str(summary.passed_tests)
        row[2].text = str(summary.failed_tests)
        row[3].text = f"{summary.success_rate:.1f}%"

        doc.add_paragraph()

    def _add_test_result(self, doc: Document, result: TestResult) -> None:
        """Add a single test result section."""
        # Test header
        status = "PASSED" if result.passed else "FAILED"
        status_color = "green" if result.passed else "red"

        doc.add_heading(f"TEST: {result.test_id}", level=2)

        # Test metadata table
        table = doc.add_table(rows=4, cols=2)
        table.style = "Table Grid"

        rows = table.rows
        rows[0].cells[0].text = "Title"
        rows[0].cells[1].text = result.title
        rows[1].cells[0].text = "Role"
        rows[1].cells[1].text = result.role
        rows[2].cells[0].text = "Status"
        rows[2].cells[1].text = status
        rows[3].cells[0].text = "Expected Result"
        rows[3].cells[1].text = result.expected_result

        # Make first column bold
        for row in rows:
            for paragraph in row.cells[0].paragraphs:
                for run in paragraph.runs:
                    run.bold = True

        doc.add_paragraph()

        # Steps
        doc.add_heading("Steps", level=3)

        for step_result in result.step_results:
            self._add_step_result(doc, step_result)

        # Verification
        doc.add_heading("Verification", level=3)

        verify_table = doc.add_table(rows=2, cols=2)
        verify_table.style = "Table Grid"

        verify_table.rows[0].cells[0].text = "Expected Result"
        verify_table.rows[0].cells[1].text = result.expected_result
        verify_table.rows[1].cells[0].text = "Observed Result"
        verify_table.rows[1].cells[1].text = result.observed_result

        if result.error:
            error_row = verify_table.add_row().cells
            error_row[0].text = "Error"
            error_row[1].text = result.error

        doc.add_paragraph()
        doc.add_paragraph("─" * 50)  # Separator
        doc.add_paragraph()

    def _add_step_result(self, doc: Document, step_result: StepResult) -> None:
        """Add a single step result with screenshot."""
        # Just add screenshot directly - no text, no timestamp
        if step_result.screenshot_path:
            screenshot_path = step_result.screenshot_path
            logger.info(f"📸 Processing screenshot: {screenshot_path}")
            logger.info(f"   Absolute path: {screenshot_path.absolute()}")
            logger.info(f"   Exists: {screenshot_path.exists()}")

            if screenshot_path.exists():
                logger.info(f"   File size: {screenshot_path.stat().st_size} bytes")
                try:
                    doc.add_picture(
                        str(screenshot_path.absolute()),  # Use absolute path
                        width=Inches(6)
                    )
                    # Center the image
                    last_paragraph = doc.paragraphs[-1]
                    last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    # Add spacing after screenshot
                    doc.add_paragraph()
                    logger.info(f"   ✅ Screenshot embedded successfully")
                except Exception as e:
                    # Only show error if screenshot failed
                    error_msg = f"[Screenshot unavailable: {e}]"
                    doc.add_paragraph(error_msg)
                    logger.error(f"   ❌ Failed to embed screenshot: {e}")
            else:
                logger.warning(f"   ⚠️ Screenshot file not found!")
                doc.add_paragraph(f"[Screenshot not found: {screenshot_path}]")
        else:
            logger.warning(f"   ⚠️ No screenshot path provided for step")

    def _generate_filename(self, summary: ExecutionSummary) -> str:
        """Generate filename for evidence document."""
        # Remove extension from workbook name
        workbook_base = Path(summary.workbook_name).stem
        date_str = summary.execution_date.strftime("%Y%m%d_%H%M%S")
        return f"CSV-Evidence-{workbook_base}-{date_str}.docx"
