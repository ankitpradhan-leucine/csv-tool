"""Excel workbook parser for test scenarios and credentials."""

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import BaseModel

from csvtool.models import Credential, TestScenario


class WorkbookData(BaseModel):
    """Parsed workbook data containing credentials and scenarios."""

    credentials: dict[str, Credential]
    scenarios: list[TestScenario]
    workbook_name: str

    def get_credential(self, role: str) -> Credential:
        """Get credential by role name."""
        if role not in self.credentials:
            raise KeyError(f"No credential found for role: {role}")
        return self.credentials[role]


class ExcelParser:
    """Parser for Excel workbooks containing test scenarios."""

    CREDENTIALS_SHEET = "Credentials"
    CREDENTIALS_HEADERS = ["Role", "Username", "Password"]
    SCENARIO_HEADERS = ["Test ID", "Role", "Title", "Test Instructions", "Expected Result"]

    def __init__(self, workbook_path: Path):
        """Initialize parser with workbook path."""
        self.workbook_path = Path(workbook_path)
        if not self.workbook_path.exists():
            raise FileNotFoundError(f"Workbook not found: {workbook_path}")

    def parse(self) -> WorkbookData:
        """Parse the workbook and return structured data."""
        wb = load_workbook(self.workbook_path, read_only=True, data_only=True)

        try:
            credentials = self._parse_credentials(wb)
            scenarios = self._parse_scenarios(wb)

            return WorkbookData(
                credentials=credentials,
                scenarios=scenarios,
                workbook_name=self.workbook_path.name
            )
        finally:
            wb.close()

    def _parse_credentials(self, wb) -> dict[str, Credential]:
        """Parse the Credentials sheet."""
        if self.CREDENTIALS_SHEET not in wb.sheetnames:
            raise ValueError(
                f"Workbook must contain a '{self.CREDENTIALS_SHEET}' sheet"
            )

        sheet = wb[self.CREDENTIALS_SHEET]
        credentials = {}

        rows = list(sheet.iter_rows(min_row=2, values_only=True))
        for row in rows:
            if row[0] is None:
                continue
            role, username, password = row[0], row[1], row[2]
            credentials[role] = Credential(
                role=str(role),
                username=str(username),
                password=str(password)
            )

        return credentials

    def _parse_scenarios(self, wb) -> list[TestScenario]:
        """Parse all scenario sheets (any sheet that's not Credentials)."""
        scenarios = []

        for sheet_name in wb.sheetnames:
            if sheet_name == self.CREDENTIALS_SHEET:
                continue

            sheet = wb[sheet_name]
            scenarios.extend(self._parse_scenario_sheet(sheet))

        return scenarios

    def _parse_scenario_sheet(self, sheet: Worksheet) -> list[TestScenario]:
        """Parse a single scenario sheet."""
        scenarios = []

        rows = list(sheet.iter_rows(min_row=2, values_only=True))
        for row in rows:
            if row[0] is None:
                continue

            test_id = str(row[0]) if row[0] else ""
            role = str(row[1]) if row[1] else ""
            title = str(row[2]) if row[2] else ""
            instructions = str(row[3]) if row[3] else ""
            expected_result = str(row[4]) if row[4] else ""

            if test_id and instructions:
                scenarios.append(TestScenario(
                    test_id=test_id,
                    role=role,
                    title=title,
                    instructions=instructions,
                    expected_result=expected_result
                ))

        return scenarios
