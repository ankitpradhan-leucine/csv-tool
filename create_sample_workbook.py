#!/usr/bin/env python3
"""
Create a sample test workbook for CSV automation tool testing.

Run this after installing dependencies:
    python3 create_sample_workbook.py
"""

from openpyxl import Workbook
from pathlib import Path


def create_sample_workbook():
    """Create sample Excel workbook with test scenarios."""

    wb = Workbook()

    # ===== Sheet 1: Credentials =====
    creds_sheet = wb.active
    creds_sheet.title = "Credentials"

    # Headers
    creds_sheet.append(["Role", "Username", "Password"])

    # Sample credentials (for testing with example.com - won't actually work)
    creds_sheet.append(["Admin", "admin@example.com", "password123"])
    creds_sheet.append(["User", "user@example.com", "userpass"])

    # ===== Sheet 2: Simple Tests =====
    test_sheet = wb.create_sheet("Simple Tests")

    # Headers
    test_sheet.append([
        "Test ID",
        "Role",
        "Title",
        "Test Instructions",
        "Expected Result"
    ])

    # Test 1: Basic navigation
    test_sheet.append([
        "TEST-001",
        "Admin",
        "Verify homepage loads",
        "1. Navigate to homepage\n2. Wait for page to load completely",
        "Homepage displays with 'Example Domain' title"
    ])

    # Test 2: Link verification
    test_sheet.append([
        "TEST-002",
        "User",
        "Verify information link",
        "1. Navigate to homepage\n2. Look for 'More information' link\n3. Verify link is visible",
        "Link is visible on the page"
    ])

    # Test 3: Page content
    test_sheet.append([
        "TEST-003",
        "Admin",
        "Verify page text content",
        "1. Navigate to homepage\n2. Read main paragraph text",
        "Page contains text about example domain usage"
    ])

    # ===== Sheet 3: Advanced Tests (for real applications) =====
    advanced_sheet = wb.create_sheet("Advanced Tests - Template")

    # Headers
    advanced_sheet.append([
        "Test ID",
        "Role",
        "Title",
        "Test Instructions",
        "Expected Result"
    ])

    # Template for real application tests
    advanced_sheet.append([
        "APP-001",
        "Admin",
        "[TEMPLATE] Login and navigate to settings",
        "1. Navigate to login page\n2. Enter credentials\n3. Click login button\n4. Navigate to settings menu\n5. Click on user settings",
        "User settings page is displayed"
    ])

    advanced_sheet.append([
        "APP-002",
        "Quality Specialist",
        "[TEMPLATE] Create new record with validation",
        "1. Navigate to records page\n2. Click 'Create New' button\n3. Fill in all required fields\n4. Leave one mandatory field blank\n5. Click Submit",
        "Validation error message is displayed"
    ])

    # ===== Add column widths for readability =====
    for sheet in wb.worksheets:
        sheet.column_dimensions['A'].width = 12  # Test ID
        sheet.column_dimensions['B'].width = 20  # Role
        sheet.column_dimensions['C'].width = 35  # Title
        sheet.column_dimensions['D'].width = 50  # Instructions
        sheet.column_dimensions['E'].width = 40  # Expected Result

    # Save workbook
    output_path = Path("scenarios/sample-test.xlsx")
    output_path.parent.mkdir(exist_ok=True)

    wb.save(output_path)
    print(f"✓ Created sample workbook: {output_path}")
    print(f"\nWorkbook contains:")
    print(f"  - Credentials sheet: 2 sample users")
    print(f"  - Simple Tests: 3 basic tests for example.com")
    print(f"  - Advanced Tests: 2 templates for real applications")
    print(f"\nYou can now run:")
    print(f"  csvtool run --url 'https://example.com' --workbook '{output_path}' --no-headless")


if __name__ == "__main__":
    try:
        create_sample_workbook()
    except ImportError as e:
        print("Error: Required dependencies not installed")
        print(f"\nPlease run: pip install -e .")
        print(f"\nMissing: {e}")
        exit(1)
    except Exception as e:
        print(f"Error creating workbook: {e}")
        exit(1)
